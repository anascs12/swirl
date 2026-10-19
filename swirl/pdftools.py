"""The PDF toolkit: every tool takes files plus options and writes new files
beside the originals. Swirl Studio builds each tool's form from its options.
"""

from __future__ import annotations

import io
import math
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from PIL import Image
from pypdf import PdfReader, PdfWriter, Transformation
from pypdf.constants import UserAccessPermissions as Perm
from pypdf.generic import RectangleObject

from . import tools
from .engine import KINDS, beside, free_path

PDF = {".pdf"}
IMAGES = KINDS["image"]


# ── Options and tools ─────────────────────────────────────────

@dataclass
class Option:
    key: str
    label: str
    type: str = "text"        # text | password | choice | int | float | pages | bool | file
    default: Any = None
    choices: tuple[str, ...] = ()
    hint: str = ""
    min: float | None = None
    max: float | None = None


@dataclass
class Tool:
    id: str
    title: str
    blurb: str
    category: str
    accepts: set[str]
    run: Callable[[list[Path], dict], list[Path]]
    options: list[Option] = field(default_factory=list)
    min_files: int = 1
    needs: tuple[str, ...] = ()
    page: str | None = None   # visual tools open an editor instead of running directly

    def accepts_file(self, path: Path) -> bool:
        return path.suffix.lower() in self.accepts

    def missing(self) -> list[str]:
        return [n for n in self.needs if not tools.find(n)]

    def defaults(self) -> dict:
        return {o.key: o.default for o in self.options}


def run_tool(tool: Tool, files: list[Path], options: dict | None = None) -> list[Path]:
    missing = tool.missing()
    if missing:
        raise tools.ToolError(f"{tool.title} needs {', '.join(missing)} installed")
    if len(files) < tool.min_files:
        raise ValueError(f"{tool.title} needs at least {tool.min_files} files")
    bad = [f.name for f in files if not tool.accepts_file(f)]
    if bad:
        raise ValueError(f"{tool.title} can't use {', '.join(bad)}")
    opts = {**tool.defaults(), **(options or {})}
    return tool.run(files, opts)


# ── Helpers ───────────────────────────────────────────────────

class PageRangeError(ValueError):
    pass


def parse_ranges(spec: str, count: int) -> list[list[int]]:
    """'1-3, 5, 8-' → [[0,1,2],[4],[7..count-1]] (0-based; 'last' or empty end allowed)."""
    spec = (spec or "").strip().lower()
    if not spec or spec == "all":
        return [list(range(count))]
    groups = []
    for part in re.split(r"[,;\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d+|last)?\s*-\s*(\d+|last)?|(\d+|last)", part)
        if not m:
            raise PageRangeError(f"Can't read page range '{part}'")
        num = lambda s, d: count if s == "last" else (int(s) if s else d)
        if m.group(3):
            a = b = num(m.group(3), 1)
        else:
            a, b = num(m.group(1), 1), num(m.group(2), count)
        if not (1 <= a <= count and 1 <= b <= count):
            raise PageRangeError(f"Page {max(a, b)} is outside 1–{count}")
        step = 1 if b >= a else -1
        groups.append(list(range(a - 1, b - 1 + step, step)))
    return groups


def flat(groups: list[list[int]]) -> list[int]:
    return list(dict.fromkeys(i for g in groups for i in g))


def open_reader(path: Path, password: str = "") -> PdfReader:
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        if not password or not reader.decrypt(password):
            raise ValueError(f"{path.name} is password-protected. Use Unlock first.")
    return reader


def write(writer: PdfWriter, dest: Path) -> Path:
    with open(dest, "wb") as fh:
        writer.write(fh)
    return dest


def overlay_page(width: float, height: float, draw: Callable) -> "PdfReader":
    """A one-page PDF of the given size, drawn with reportlab, for stamping."""
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    draw(c, width, height)
    c.save()
    buf.seek(0)
    return PdfReader(buf)


def _hex(color: str):
    from reportlab.lib.colors import HexColor
    try:
        return HexColor(color if color.startswith("#") else "#" + color)
    except Exception:
        return HexColor("#555555")


# ── Organize ──────────────────────────────────────────────────

def merge(files, o):
    writer = PdfWriter()
    for f in files:
        writer.append(open_reader(f))
    return [write(writer, beside(files[0], ".pdf", "-merged"))]


def split(files, o):
    out = []
    for src in files:
        reader = open_reader(src)
        n = len(reader.pages)
        if o["mode"] == "Every page":
            groups = [[i] for i in range(n)]
        elif o["mode"] == "Every N pages":
            step = max(1, int(o["every"]))
            groups = [list(range(i, min(i + step, n))) for i in range(0, n, step)]
        else:
            groups = parse_ranges(o["ranges"], n)
        for g in groups:
            writer = PdfWriter()
            for i in g:
                writer.add_page(reader.pages[i])
            label = f"{g[0] + 1}" if len(g) == 1 else f"{g[0] + 1}-{g[-1] + 1}"
            out.append(write(writer, beside(src, ".pdf", f"-pages {label}")))
    return out


def extract_pages(files, o):
    out = []
    for src in files:
        reader = open_reader(src)
        writer = PdfWriter()
        for i in flat(parse_ranges(o["pages"], len(reader.pages))):
            writer.add_page(reader.pages[i])
        out.append(write(writer, beside(src, ".pdf", "-extract")))
    return out


def remove_pages(files, o):
    out = []
    for src in files:
        reader = open_reader(src)
        drop = set(flat(parse_ranges(o["pages"], len(reader.pages))))
        keep = [i for i in range(len(reader.pages)) if i not in drop]
        if not keep:
            raise ValueError("That would remove every page")
        writer = PdfWriter()
        for i in keep:
            writer.add_page(reader.pages[i])
        out.append(write(writer, beside(src, ".pdf", "-edited")))
    return out


def reorder(files, o):
    """New page order, e.g. '3,1,2' or 'reverse'."""
    out = []
    for src in files:
        reader = open_reader(src)
        n = len(reader.pages)
        order = list(range(n - 1, -1, -1)) if o["order"].strip().lower() == "reverse" \
            else [i for g in parse_ranges(o["order"], n) for i in g]
        writer = PdfWriter()
        for i in order:
            writer.add_page(reader.pages[i])
        out.append(write(writer, beside(src, ".pdf", "-reordered")))
    return out


def rotate(files, o):
    angle = {"90° right": 90, "180°": 180, "90° left": 270}[o["angle"]]
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        for i in flat(parse_ranges(o["pages"], len(writer.pages))):
            writer.pages[i].rotate(angle)
        out.append(write(writer, beside(src, ".pdf", "-rotated")))
    return out


# ── Optimize ──────────────────────────────────────────────────

def compress(files, o):
    level = o["level"]
    quality, max_side = {"Light": (85, None), "Recommended": (65, 2000), "Extreme": (40, 1200)}[level]
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        for page in writer.pages:
            for img in page.images:
                try:
                    pil = img.image
                    if max_side and max(pil.size) > max_side:
                        pil = pil.copy()
                        pil.thumbnail((max_side, max_side))
                    if pil.mode not in ("RGB", "L"):
                        pil = pil.convert("RGB")
                    img.replace(pil, quality=quality)
                except Exception:
                    pass  # leave images pypdf cannot re-encode untouched
            page.compress_content_streams()
        writer.compress_identical_objects()
        dest = beside(src, ".pdf", "-compressed")
        write(writer, dest)
        if dest.stat().st_size >= src.stat().st_size:  # never hand back a bigger file
            dest.unlink()
            shutil.copyfile(src, dest)
        out.append(dest)
    return out


def repair(files, o):
    out = []
    for src in files:
        reader = PdfReader(src, strict=False)
        writer = PdfWriter()
        for page in reader.pages:
            try:
                writer.add_page(page)
            except Exception:
                continue  # skip pages too broken to copy
        if not writer.pages:
            raise ValueError(f"No readable pages in {src.name}")
        out.append(write(writer, beside(src, ".pdf", "-repaired")))
    return out


def crop(files, o):
    """Trim margins: 'Auto' finds the printed area of each page."""
    import pypdfium2 as pdfium
    out = []
    pad = float(o["padding"]) * 72 / 25.4
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        doc = pdfium.PdfDocument(src) if o["mode"] == "Auto (trim white space)" else None
        for i, page in enumerate(writer.pages):
            box = page.cropbox
            x0, y0, x1, y1 = float(box.left), float(box.bottom), float(box.right), float(box.top)
            if doc is not None:
                scale = 1.0
                img = doc[i].render(scale=scale).to_pil().convert("L")
                bbox = Image.eval(img, lambda v: 255 if v < 245 else 0).getbbox()
                if not bbox:
                    continue
                left, top, right, bottom = bbox
                h = img.height
                nx0, nx1 = x0 + left / scale - pad, x0 + right / scale + pad
                ny0, ny1 = y0 + (h - bottom) / scale - pad, y0 + (h - top) / scale + pad
            else:
                m = float(o["margin"]) * 72 / 25.4
                nx0, ny0, nx1, ny1 = x0 + m, y0 + m, x1 - m, y1 - m
            nx0, ny0, nx1, ny1 = max(nx0, x0), max(ny0, y0), min(nx1, x1), min(ny1, y1)
            if nx1 - nx0 > 10 and ny1 - ny0 > 10:
                page.cropbox = RectangleObject((nx0, ny0, nx1, ny1))
                page.mediabox = RectangleObject((nx0, ny0, nx1, ny1))
        if doc is not None:
            doc.close()
        out.append(write(writer, beside(src, ".pdf", "-cropped")))
    return out


# ── Edit ──────────────────────────────────────────────────────

_POSITIONS = ("Bottom center", "Bottom right", "Bottom left", "Top center", "Top right", "Top left")


def page_numbers(files, o):
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        targets = set(flat(parse_ranges(o["pages"], len(writer.pages))))
        total = len(targets)
        n = int(o["start"])
        for i, page in enumerate(writer.pages):
            if i not in targets:
                continue
            text = o["format"].replace("{n}", str(n)).replace("{total}", str(total + int(o["start"]) - 1))
            n += 1
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            pos, size, margin = o["position"], float(o["size"]), 28

            def draw(c, w, h, text=text):
                c.setFont("Helvetica", size)
                c.setFillColor(_hex(o["color"]))
                y = margin if pos.startswith("Bottom") else h - margin - size
                if pos.endswith("center"):
                    c.drawCentredString(w / 2, y, text)
                elif pos.endswith("right"):
                    c.drawRightString(w - margin, y, text)
                else:
                    c.drawString(margin, y, text)

            stamp = overlay_page(w, h, draw).pages[0]
            page.merge_transformed_page(stamp, Transformation().translate(float(page.mediabox.left), float(page.mediabox.bottom)))
        out.append(write(writer, beside(src, ".pdf", "-numbered")))
    return out


def watermark(files, o):
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        targets = set(flat(parse_ranges(o["pages"], len(writer.pages))))
        image = Path(o["image"]) if o.get("image") else None
        for i, page in enumerate(writer.pages):
            if i not in targets:
                continue
            w, h = float(page.mediabox.width), float(page.mediabox.height)

            def draw(c, w, h):
                c.saveState()
                c.setFillAlpha(float(o["opacity"]) / 100)
                if image:
                    from reportlab.lib.utils import ImageReader
                    ir = ImageReader(str(image))
                    iw, ih = ir.getSize()
                    scale = min(w * 0.6 / iw, h * 0.6 / ih)
                    c.translate(w / 2, h / 2)
                    c.rotate(float(o["angle"]))
                    c.drawImage(ir, -iw * scale / 2, -ih * scale / 2, iw * scale, ih * scale, mask="auto")
                else:
                    size = float(o["size"])
                    c.setFont("Helvetica-Bold", size)
                    c.setFillColor(_hex(o["color"]), alpha=float(o["opacity"]) / 100)
                    if o["layout"] == "Tiled":
                        step_x, step_y = max(size * len(o["text"]) * 0.7, 120), size * 4
                        for gx in range(-1, int(w // step_x) + 2):
                            for gy in range(-1, int(h // step_y) + 2):
                                c.saveState()
                                c.translate(gx * step_x + (gy % 2) * step_x / 2, gy * step_y)
                                c.rotate(float(o["angle"]))
                                c.drawCentredString(0, 0, o["text"])
                                c.restoreState()
                    else:
                        c.translate(w / 2, h / 2)
                        c.rotate(float(o["angle"]))
                        c.drawCentredString(0, -size / 3, o["text"])
                c.restoreState()

            stamp = overlay_page(w, h, draw).pages[0]
            page.merge_transformed_page(stamp, Transformation().translate(float(page.mediabox.left), float(page.mediabox.bottom)),
                                        over=o["placement"] == "Over content")
        out.append(write(writer, beside(src, ".pdf", "-watermarked")))
    return out


def edit_metadata(files, o):
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        if o["remove_all"]:
            writer.metadata = None
            if "/Metadata" in writer.root_object:
                del writer.root_object["/Metadata"]
        else:
            meta = {f"/{k.title()}": o[k] for k in ("title", "author", "subject", "keywords") if o.get(k)}
            writer.add_metadata(meta)
        out.append(write(writer, beside(src, ".pdf", "-info")))
    return out


# ── Security ──────────────────────────────────────────────────

def protect(files, o):
    if not o["password"]:
        raise ValueError("Enter a password")
    if o["password"] != o["confirm"]:
        raise ValueError("The two passwords don't match")
    perms = Perm.all()
    if not o["allow_print"]:
        perms &= ~(Perm.PRINT | Perm.PRINT_TO_REPRESENTATION)
    if not o["allow_copy"]:
        perms &= ~Perm.EXTRACT_TEXT_AND_GRAPHICS & ~Perm.EXTRACT
    if not o["allow_edit"]:
        perms &= ~Perm.MODIFY & ~Perm.ADD_OR_MODIFY & ~Perm.ASSEMBLE_DOC & ~Perm.FILL_FORM_FIELDS
    out = []
    for src in files:
        writer = PdfWriter(clone_from=open_reader(src))
        writer.encrypt(user_password=o["password"], owner_password=None, permissions_flag=perms, algorithm="AES-256")
        out.append(write(writer, beside(src, ".pdf", "-protected")))
    return out


def unlock(files, o):
    out = []
    for src in files:
        reader = PdfReader(src, strict=False)
        if reader.is_encrypted and not reader.decrypt(o["password"] or ""):
            raise ValueError(f"Wrong password for {src.name}")
        writer = PdfWriter(clone_from=reader)
        out.append(write(writer, beside(src, ".pdf", "-unlocked")))
    return out


# ── Convert from PDF ──────────────────────────────────────────

def pdf_to_images(files, o):
    import pypdfium2 as pdfium
    fmt = o["format"].lower()
    suffix = ".jpg" if fmt == "jpg" else ".png"
    scale = int(o["dpi"]) / 72
    out = []
    for src in files:
        doc = pdfium.PdfDocument(src, password=None)
        n = len(doc)
        for i in flat(parse_ranges(o["pages"], n)):
            img = doc[i].render(scale=scale).to_pil().convert("RGB")
            dest = beside(src, suffix, f"-page{i + 1}" if n > 1 else "")
            img.save(dest, "JPEG" if suffix == ".jpg" else "PNG", **({"quality": 90} if suffix == ".jpg" else {}))
            out.append(dest)
        doc.close()
    return out


def pdf_to_text(files, o):
    out = []
    for src in files:
        reader = open_reader(src)
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        dest = beside(src, ".txt")
        dest.write_text(text, encoding="utf-8")
        out.append(dest)
    return out


def extract_images(files, o):
    out = []
    for src in files:
        reader = open_reader(src)
        folder = free_path(src.parent, src.stem + "-images", "")
        count = 0
        for pi, page in enumerate(reader.pages):
            for ii, img in enumerate(page.images):
                folder.mkdir(exist_ok=True)
                ext = Path(img.name).suffix or ".png"
                (folder / f"page{pi + 1}-{ii + 1}{ext}").write_bytes(img.data)
                count += 1
        if not count:
            raise ValueError(f"No images found in {src.name}")
        out.append(folder)
    return out


def pdf_to_pptx(files, o):
    """One slide per page, the page as a sharp image (layout stays exact)."""
    import pypdfium2 as pdfium
    from pptx import Presentation
    from pptx.util import Emu
    out = []
    for src in files:
        doc = pdfium.PdfDocument(src)
        prs = Presentation()
        w0, h0 = doc[0].get_size()
        prs.slide_width = Emu(int(w0 * 12700))
        prs.slide_height = Emu(int(h0 * 12700))
        with tempfile.TemporaryDirectory() as tmp:
            for i in range(len(doc)):
                png = Path(tmp, f"{i}.png")
                doc[i].render(scale=2).to_pil().convert("RGB").save(png)
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        doc.close()
        dest = beside(src, ".pptx")
        prs.save(dest)
        out.append(dest)
    return out


def _soffice(src: Path, fmt: str, suffix: str, infilter: str | None = None) -> Path:
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "profile").resolve().as_uri()
        args = [tools.find("soffice"), f"-env:UserInstallation={profile}", "--headless"]
        if infilter:
            args.append(f"--infilter={infilter}")
        tools.run([*args, "--convert-to", fmt, "--outdir", tmp, str(src)], timeout=300)
        produced = Path(tmp, src.stem + suffix)
        if not produced.exists():
            raise tools.ToolError(f"LibreOffice could not convert {src.name}")
        dest = beside(src, suffix)
        shutil.move(str(produced), dest)
        return dest


def pdf_to_word(files, o):
    return [_soffice(f, "docx:MS Word 2007 XML", ".docx", infilter="writer_pdf_import") for f in files]


def pdf_to_excel(files, o):
    """Tables found on each page become sheets; pages without ruled tables fall back to text lines."""
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    out = []
    for src in files:
        wb = Workbook()
        wb.remove(wb.active)
        with pdfplumber.open(src) as pdf:
            for pi, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                if not tables and o["fallback"]:
                    words = page.extract_text_lines() or []
                    tables = [[[w["text"]] for w in words]] if words else []
                for ti, table in enumerate(tables, start=1):
                    ws = wb.create_sheet(f"Page {pi}" + (f" ({ti})" if len(tables) > 1 else ""))
                    for row in table:
                        ws.append([_number(c) for c in row])
                    for ci, col in enumerate(ws.columns, start=1):
                        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                        ws.column_dimensions[get_column_letter(ci)].width = min(60, width + 2)
        if not wb.sheetnames:
            raise ValueError(f"No tables found in {src.name}")
        dest = beside(src, ".xlsx")
        wb.save(dest)
        out.append(dest)
    return out


def _number(cell):
    """'1,234.50' → 1234.5 so Excel can sum it; anything else stays text."""
    if cell is None:
        return None
    text = str(cell).strip()
    plain = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", plain):
        return float(plain) if "." in plain else int(plain)
    return text


def pdf_to_markdown(files, o):
    """Text with headings guessed from font size, and tables as Markdown tables."""
    import pdfplumber
    out = []
    for src in files:
        parts = []
        with pdfplumber.open(src) as pdf:
            sizes = [c["size"] for p in pdf.pages[:20] for c in p.chars]
            body = sorted(sizes)[len(sizes) // 2] if sizes else 10
            for page in pdf.pages:
                table_boxes = [t.bbox for t in page.find_tables()]
                for t in page.extract_tables():
                    rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in t]
                    if rows:
                        parts.append("| " + " | ".join(rows[0]) + " |")
                        parts.append("|" + "---|" * len(rows[0]))
                        parts += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                        parts.append("")
                outside = page
                for box in table_boxes:
                    outside = outside.outside_bbox(box)
                for line in outside.extract_text_lines() or []:
                    size = max((c["size"] for c in line["chars"]), default=body)
                    text = line["text"].strip()
                    if not text:
                        continue
                    if size >= body * 1.6:
                        parts.append(f"\n# {text}\n")
                    elif size >= body * 1.25:
                        parts.append(f"\n## {text}\n")
                    else:
                        parts.append(text)
                parts.append("")
        dest = beside(src, ".md")
        dest.write_text(re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip() + "\n", encoding="utf-8")
        out.append(dest)
    return out


def pdf_to_pdfa(files, o):
    """PDF/A-2b for archiving: Ghostscript when installed, otherwise LibreOffice."""
    out = []
    for src in files:
        dest = beside(src, ".pdf", "-pdfa")
        if tools.find("gs"):
            tools.run([tools.find("gs"), "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE",
                       "-sColorConversionStrategy=RGB", "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1",
                       f"-sOutputFile={dest}", str(src)])
        else:
            options = '{"SelectPdfVersion":{"type":"long","value":"2"}}'
            made = _soffice(src, f"pdf:draw_pdf_Export:{options}", ".pdf", infilter="draw_pdf_import")
            made.rename(dest)
        out.append(dest)
    return out


def compare(files, o):
    """Side-by-side report of what changed between two PDFs (text and pages)."""
    import difflib
    import html
    import base64
    import pypdfium2 as pdfium
    from PIL import ImageChops

    a, b = files[0], files[1]
    ra, rb = open_reader(a), open_reader(b)
    ta = [(p.extract_text() or "").splitlines() for p in ra.pages]
    tb = [(p.extract_text() or "").splitlines() for p in rb.pages]
    rows, changed_pages = [], []
    da, db = pdfium.PdfDocument(a), pdfium.PdfDocument(b)
    for i in range(max(len(ta), len(tb))):
        la = ta[i] if i < len(ta) else []
        lb = tb[i] if i < len(tb) else []
        diff = list(difflib.ndiff(la, lb))
        changes = [d for d in diff if d[:2] in ("- ", "+ ")]
        pixel_change = False
        if i < len(da) and i < len(db):
            ia = da[i].render(scale=0.5).to_pil().convert("L")
            ib = db[i].render(scale=0.5).to_pil().convert("L").resize(ia.size)
            pixel_change = ImageChops.difference(ia, ib).getbbox() is not None
        if changes or pixel_change or i >= len(ta) or i >= len(tb):
            changed_pages.append(i + 1)
            lines = []
            for d in diff:
                cls = {"- ": "del", "+ ": "ins"}.get(d[:2])
                if cls:
                    lines.append(f'<div class="{cls}">{html.escape(d[2:])}</div>')
            note = "page only in first file" if i >= len(tb) else "page only in second file" if i >= len(ta) \
                else ("layout or images changed" if pixel_change and not changes else "")
            thumbs = ""
            if i < len(da) and i < len(db) and o["thumbnails"]:
                pics = []
                for doc in (da, db):
                    buf = io.BytesIO()
                    doc[i].render(scale=0.6).to_pil().convert("RGB").save(buf, "JPEG", quality=70)
                    pics.append(base64.b64encode(buf.getvalue()).decode())
                thumbs = f'<div class="thumbs"><img src="data:image/jpeg;base64,{pics[0]}"><img src="data:image/jpeg;base64,{pics[1]}"></div>'
            rows.append(f"<section><h2>Page {i + 1}</h2><p class=note>{note}</p>{thumbs}{''.join(lines)}</section>")
    da.close(); db.close()
    summary = (f"{len(changed_pages)} of {max(len(ta), len(tb))} pages differ: {', '.join(map(str, changed_pages))}"
               if changed_pages else "No differences found.")
    page = f"""<!doctype html><meta charset=utf-8><title>Compare {html.escape(a.name)} and {html.escape(b.name)}</title>
<style>body{{font:15px/1.5 system-ui,sans-serif;max-width:1000px;margin:32px auto;padding:0 16px;color:#111}}
h1{{font-size:22px}}.del{{background:#fde8e8;color:#9b1c1c;padding:2px 8px;border-radius:4px;margin:2px 0}}
.del::before{{content:"− "}}.ins{{background:#def7ec;color:#03543f;padding:2px 8px;border-radius:4px;margin:2px 0}}
.ins::before{{content:"+ "}}section{{border-top:1px solid #e5e7eb;padding:8px 0}}.note{{color:#6b7280}}
.thumbs{{display:flex;gap:12px}}.thumbs img{{width:48%;border:1px solid #e5e7eb}}</style>
<h1>{html.escape(a.name)} → {html.escape(b.name)}</h1><p>{summary}</p>{''.join(rows)}"""
    dest = beside(a, ".html", "-compare")
    dest.write_text(page, encoding="utf-8")
    return [dest]


# ── Convert to PDF ────────────────────────────────────────────

_PAGE_SIZES = {"A4": (595.28, 841.89), "Letter": (612, 792)}


def images_to_pdf(files, o):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import ImageOps
    margin = float(o["margin"]) * 72 / 25.4
    dest = beside(files[0], ".pdf", "-images" if len(files) > 1 else "")
    c = canvas.Canvas(str(dest))
    for f in files:
        img = ImageOps.exif_transpose(Image.open(f)).convert("RGB")
        iw, ih = img.size
        if o["page_size"] == "Fit to image":
            pw, ph = iw * 72 / 150 + 2 * margin, ih * 72 / 150 + 2 * margin
        else:
            pw, ph = _PAGE_SIZES[o["page_size"]]
            landscape = o["orientation"] == "Landscape" or (o["orientation"] == "Auto" and iw > ih)
            if landscape:
                pw, ph = ph, pw
        c.setPageSize((pw, ph))
        scale = min((pw - 2 * margin) / iw, (ph - 2 * margin) / ih)
        dw, dh = iw * scale, ih * scale
        c.drawImage(ImageReader(img), (pw - dw) / 2, (ph - dh) / 2, dw, dh)
        c.showPage()
    c.save()
    return [dest]


def office_to_pdf(files, o):
    return [_soffice(f, "pdf", ".pdf") for f in files]


def _visual(files, o):
    raise RuntimeError("This tool opens in the editor")


def ocr_tool(files, o):
    from .pdfedit import ocr
    return [ocr(f, language=o["language"]) for f in files]


# ── Registry ──────────────────────────────────────────────────

PAGES_ALL = Option("pages", "Pages", "pages", "all", hint="e.g. 1-3, 5, 8-last · 'all'")

TOOLS: list[Tool] = [
    # Organize
    Tool("merge", "Merge PDF", "Combine PDFs in the order you choose.", "Organize", PDF, merge, min_files=2),
    Tool("split", "Split PDF", "Split by ranges, every page, or every N pages.", "Organize", PDF, split, [
        Option("mode", "Split", "choice", "Custom ranges", ("Custom ranges", "Every page", "Every N pages")),
        Option("ranges", "Ranges", "pages", "1-2, 3-last", hint="each range becomes its own PDF"),
        Option("every", "N", "int", 2, min=1, max=1000)]),
    Tool("remove", "Remove pages", "Delete the pages you don't need.", "Organize", PDF, remove_pages, [
        Option("pages", "Pages to remove", "pages", "1", hint="e.g. 2, 5-7")]),
    Tool("extract", "Extract pages", "Keep only selected pages, in a new PDF.", "Organize", PDF, extract_pages, [
        Option("pages", "Pages to keep", "pages", "1-3")]),
    Tool("reorder", "Reorder pages", "Put pages in a new order, or reverse them.", "Organize", PDF, reorder, [
        Option("order", "New order", "pages", "reverse", hint="e.g. 3, 1, 2, 4-last · 'reverse'")]),
    Tool("rotate", "Rotate PDF", "Turn all or some pages.", "Organize", PDF, rotate, [
        Option("angle", "Rotate", "choice", "90° right", ("90° right", "180°", "90° left")), PAGES_ALL]),
    # Optimize
    Tool("compress", "Compress PDF", "Make PDFs smaller for email and upload.", "Optimize", PDF, compress, [
        Option("level", "Level", "choice", "Recommended", ("Light", "Recommended", "Extreme"))]),
    Tool("repair", "Repair PDF", "Rebuild a damaged PDF from the pages that still read.", "Optimize", PDF, repair),
    Tool("crop", "Crop PDF", "Trim white margins automatically or by a fixed amount.", "Optimize", PDF, crop, [
        Option("mode", "Mode", "choice", "Auto (trim white space)", ("Auto (trim white space)", "Fixed margin")),
        Option("padding", "Padding (mm)", "float", 4, min=0, max=50),
        Option("margin", "Fixed margin (mm)", "float", 10, min=0, max=100)]),
    # Edit
    Tool("numbers", "Page numbers", "Add page numbers in the position and style you want.", "Edit", PDF, page_numbers, [
        Option("position", "Position", "choice", "Bottom center", _POSITIONS),
        Option("format", "Format", "text", "{n}", hint="{n} = number, {total} = page count, e.g. 'Page {n} of {total}'"),
        Option("start", "Start at", "int", 1, min=0, max=100000),
        Option("size", "Font size", "int", 11, min=6, max=48),
        Option("color", "Color", "text", "#444444"),
        PAGES_ALL]),
    Tool("watermark", "Add watermark", "Stamp text or an image over your pages.", "Edit", PDF, watermark, [
        Option("text", "Text", "text", "CONFIDENTIAL"),
        Option("image", "Or an image", "file", "", hint="PNG with transparency works best"),
        Option("layout", "Layout", "choice", "Center", ("Center", "Tiled")),
        Option("size", "Font size", "int", 60, min=8, max=200),
        Option("angle", "Angle", "int", 45, min=-180, max=180),
        Option("opacity", "Opacity %", "int", 20, min=1, max=100),
        Option("color", "Color", "text", "#d32f2f"),
        Option("placement", "Placement", "choice", "Over content", ("Over content", "Behind content")),
        PAGES_ALL]),
    Tool("edit", "Edit PDF", "Add text, highlights, boxes, drawings and images.", "Edit", PDF, _visual, page="edit"),
    Tool("sign", "Sign PDF", "Draw, type or upload your signature and place it.", "Edit", PDF, _visual, page="sign"),
    Tool("organize", "Organize PDF", "Drag pages to reorder, rotate, delete or add blank pages.", "Organize", PDF, _visual, page="organize"),
    Tool("forms", "Fill PDF forms", "Fill in a PDF form's fields and optionally lock them.", "Edit", PDF, _visual, page="forms"),
    Tool("metadata", "Edit metadata", "Change title and author, or remove all hidden info.", "Edit", PDF, edit_metadata, [
        Option("remove_all", "Remove all metadata", "bool", False),
        Option("title", "Title", "text", ""), Option("author", "Author", "text", ""),
        Option("subject", "Subject", "text", ""), Option("keywords", "Keywords", "text", "")]),
    # Security
    Tool("protect", "Protect PDF", "Lock with a password (AES-256) and limit printing or copying.", "Security", PDF, protect, [
        Option("password", "Password", "password", ""), Option("confirm", "Repeat password", "password", ""),
        Option("allow_print", "Allow printing", "bool", True),
        Option("allow_copy", "Allow copying text", "bool", False),
        Option("allow_edit", "Allow editing", "bool", False)]),
    Tool("redact", "Redact PDF", "Permanently black out text and images.", "Security", PDF, _visual, page="redact"),
    Tool("unlock", "Unlock PDF", "Remove the password from a PDF you can open.", "Security", PDF, unlock, [
        Option("password", "Password", "password", "")]),
    # Convert from PDF
    Tool("pdf2img", "PDF to JPG / PNG", "Turn each page into an image.", "From PDF", PDF, pdf_to_images, [
        Option("format", "Format", "choice", "JPG", ("JPG", "PNG")),
        Option("dpi", "Quality (DPI)", "choice", "150", ("72", "150", "300")), PAGES_ALL]),
    Tool("pdf2word", "PDF to Word", "Editable DOCX (best for text-based PDFs).", "From PDF", PDF, pdf_to_word, needs=("soffice",)),
    Tool("pdf2pptx", "PDF to PowerPoint", "One slide per page, layout kept exactly.", "From PDF", PDF, pdf_to_pptx),
    Tool("pdf2excel", "PDF to Excel", "Tables become spreadsheet sheets, numbers stay numbers.", "From PDF", PDF, pdf_to_excel, [
        Option("fallback", "Include text lines when a page has no table", "bool", False)]),
    Tool("pdf2md", "PDF to Markdown", "Clean text with headings and tables, ready for notes or AI.", "From PDF", PDF, pdf_to_markdown),
    Tool("pdfa", "PDF to PDF/A", "Archive-safe PDF/A-2b for long-term storage.", "From PDF", PDF, pdf_to_pdfa, needs=("soffice",)),
    Tool("ocr", "OCR PDF", "Make scanned PDFs searchable and selectable.", "Optimize", PDF, ocr_tool, [
        Option("language", "Language", "choice", "eng", ("eng", "urd", "ara", "fra", "deu", "spa", "hin"),
               hint="Extra languages need their Tesseract language pack")], needs=("tesseract",)),
    Tool("pdf2txt", "PDF to text", "Extract all the text.", "From PDF", PDF, pdf_to_text),
    Tool("pdfimages", "Extract images", "Save every image inside the PDF.", "From PDF", PDF, extract_images),
    # Convert to PDF
    Tool("img2pdf", "JPG to PDF", "Images to a PDF, one per page.", "To PDF", IMAGES, images_to_pdf, [
        Option("page_size", "Page size", "choice", "A4", ("A4", "Letter", "Fit to image")),
        Option("orientation", "Orientation", "choice", "Auto", ("Auto", "Portrait", "Landscape")),
        Option("margin", "Margin (mm)", "float", 10, min=0, max=50)]),
    Tool("word2pdf", "Word to PDF", "DOC, DOCX, ODT and RTF to PDF.", "To PDF", {".doc", ".docx", ".odt", ".rtf", ".txt"}, office_to_pdf, needs=("soffice",)),
    Tool("ppt2pdf", "PowerPoint to PDF", "PPT, PPTX and ODP slideshows to PDF.", "To PDF", {".ppt", ".pptx", ".odp"}, office_to_pdf, needs=("soffice",)),
    Tool("excel2pdf", "Excel to PDF", "XLS, XLSX and ODS spreadsheets to PDF.", "To PDF", {".xls", ".xlsx", ".ods", ".csv"}, office_to_pdf, needs=("soffice",)),
    Tool("html2pdf", "HTML to PDF", "Saved web pages and HTML files to PDF.", "To PDF", {".html", ".htm"}, office_to_pdf, needs=("soffice",)),
    # Review
    Tool("compare", "Compare PDF", "See what changed between two versions, page by page.", "Review", PDF, compare, [
        Option("thumbnails", "Show page previews", "bool", True)], min_files=2),
]

TOOL_BY_ID = {t.id: t for t in TOOLS}
CATEGORIES = list(dict.fromkeys(t.category for t in TOOLS))
