from pathlib import Path

import pytest
from PIL import Image
from pypdf import PdfReader

from swirl import pdftools as pt
from swirl import tools

needs_office = pytest.mark.skipif(not tools.find("soffice"), reason="LibreOffice not installed")


def make_report(path: Path, pages: int = 4, title: str = "Quarterly Report", body_word: str = "revenue") -> Path:
    """A text PDF with a heading, a ruled table on page 1 and a photo on page 2."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    photo = path.with_suffix(".jpg")
    Image.new("RGB", (1600, 1200), (30, 120, 200)).save(photo, quality=95)
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(pages):
        c.setFont("Helvetica-Bold", 24)
        c.drawString(72, 760, f"{title} {i + 1}")
        c.setFont("Helvetica", 11)
        c.drawString(72, 730, f"Page {i + 1} talks about {body_word} and growth.")
        if i == 0:
            rows = [("Region", "Sales"), ("North", "1,200"), ("South", "950")]
            for r, (a, b) in enumerate(rows):
                y = 680 - r * 20
                c.drawString(80, y + 6, a)
                c.drawString(230, y + 6, b)
            for r in range(len(rows) + 1):
                c.line(72, 700 - r * 20, 372, 700 - r * 20)
            for x in (72, 222, 372):
                c.line(x, 700, x, 700 - 20 * len(rows))
        if i == 1:
            c.drawImage(str(photo), 72, 300, 400, 300)
        c.showPage()
    c.save()
    return path


@pytest.fixture
def report(tmp_path):
    return make_report(tmp_path / "report.pdf")


def run(tool_id, files, **opts):
    return pt.run_tool(pt.TOOL_BY_ID[tool_id], files, opts)


def pages(path):
    return len(PdfReader(path).pages)


def text(path, i=0):
    return PdfReader(path).pages[i].extract_text()


def test_parse_ranges():
    assert pt.parse_ranges("1-3, 5, 8-", 10) == [[0, 1, 2], [4], [7, 8, 9]]
    assert pt.parse_ranges("all", 3) == [[0, 1, 2]]
    assert pt.parse_ranges("last", 4) == [[3]]
    assert pt.parse_ranges("3-1", 4) == [[2, 1, 0]]
    with pytest.raises(pt.PageRangeError):
        pt.parse_ranges("9", 4)
    with pytest.raises(pt.PageRangeError):
        pt.parse_ranges("abc", 4)


def test_organize(report, tmp_path):
    other = make_report(tmp_path / "appendix.pdf", 2, title="Appendix")
    merged = run("merge", [report, other])[0]
    assert pages(merged) == 6 and "Appendix" in text(merged, 4)

    parts = run("split", [report], mode="Custom ranges", ranges="1-2, 3-last")
    assert [pages(p) for p in parts] == [2, 2]
    assert len(run("split", [report], mode="Every page")) == 4
    assert [pages(p) for p in run("split", [report], mode="Every N pages", every=3)] == [3, 1]

    assert pages(run("remove", [report], pages="2-3")[0]) == 2
    kept = run("extract", [report], pages="4, 1")[0]
    assert "Report 4" in text(kept, 0) and "Report 1" in text(kept, 1)
    rev = run("reorder", [report], order="reverse")[0]
    assert "Report 4" in text(rev, 0)
    rotated = run("rotate", [report], angle="90° right", pages="1")[0]
    assert PdfReader(rotated).pages[0].rotation == 90 and PdfReader(rotated).pages[1].rotation == 0
    with pytest.raises(ValueError, match="every page"):
        run("remove", [report], pages="all")


def test_compress_never_grows(report):
    for level in ("Light", "Recommended", "Extreme"):
        out = run("compress", [report], level=level)[0]
        assert out.stat().st_size <= report.stat().st_size
        assert pages(out) == 4
    extreme = run("compress", [report], level="Extreme")[0]
    assert extreme.stat().st_size < report.stat().st_size


def test_edit_tools(report):
    numbered = run("numbers", [report], format="Page {n} of {total}")[0]
    assert "Page 1 of 4" in text(numbered, 0) and "Page 4 of 4" in text(numbered, 3)
    marked = run("watermark", [report], text="DRAFT", layout="Tiled")[0]
    assert "DRAFT" in text(marked, 2)
    info = run("metadata", [report], title="Final", author="Ayesha")[0]
    assert PdfReader(info).metadata.title == "Final"
    wiped = run("metadata", [info], remove_all=True)[0]
    assert not (PdfReader(wiped).metadata or {}).get("/Author")


def test_crop_auto_trims_white_space(report):
    out = run("crop", [report], mode="Auto (trim white space)", padding=2)[0]
    before, after = PdfReader(report).pages[0].mediabox, PdfReader(out).pages[0].mediabox
    assert after.width < before.width and after.height < before.height


def test_protect_and_unlock(report):
    with pytest.raises(ValueError, match="match"):
        run("protect", [report], password="a", confirm="b")
    locked = run("protect", [report], password="s3cret", confirm="s3cret")[0]
    reader = PdfReader(locked)
    assert reader.is_encrypted
    with pytest.raises(ValueError, match="password-protected"):
        run("merge", [locked, report])
    with pytest.raises(ValueError, match="Wrong password"):
        run("unlock", [locked], password="nope")
    opened = run("unlock", [locked], password="s3cret")[0]
    assert not PdfReader(opened).is_encrypted and pages(opened) == 4


def test_from_pdf(report):
    jpgs = run("pdf2img", [report], format="JPG", dpi="72", pages="1-2")
    assert [p.name for p in jpgs] == ["report-page1.jpg", "report-page2.jpg"]
    assert "growth" in run("pdf2txt", [report])[0].read_text()
    folder = run("pdfimages", [report])[0]
    assert len(list(folder.iterdir())) == 1
    pptx = run("pdf2pptx", [report])[0]
    from pptx import Presentation
    assert len(Presentation(pptx).slides) == 4

    xlsx = run("pdf2excel", [report])[0]
    from openpyxl import load_workbook
    ws = load_workbook(xlsx)["Page 1"]
    assert [c.value for c in ws[2]] == ["North", 1200]       # numbers stay numbers

    md = run("pdf2md", [report])[0].read_text()
    assert "# Quarterly Report 1" in md and "| North | 1,200 |" in md


def test_images_to_pdf(tmp_path):
    a = tmp_path / "wide.png"; Image.new("RGB", (1200, 600), "red").save(a)
    b = tmp_path / "tall.jpg"; Image.new("RGB", (600, 1200), "blue").save(b)
    out = run("img2pdf", [a, b], page_size="A4", orientation="Auto")[0]
    p = PdfReader(out).pages
    assert len(p) == 2 and p[0].mediabox.width > p[0].mediabox.height   # landscape for the wide image


def test_compare(report, tmp_path):
    changed = make_report(tmp_path / "report-v2.pdf", body_word="profit")
    html = run("compare", [report, changed], thumbnails=False)[0].read_text()
    assert "4 of 4 pages differ" in html and "profit" in html
    same = run("compare", [report, report], thumbnails=False)[0].read_text()
    assert "No differences found" in same


@needs_office
def test_office_round_trip(report, tmp_path):
    docx = run("pdf2word", [report])[0]
    assert docx.suffix == ".docx"
    back = run("word2pdf", [docx])[0]
    assert "Quarterly Report" in "".join(text(back, i) for i in range(pages(back)))
    html = tmp_path / "page.html"
    html.write_text("<h1>Hello HTML</h1><p>Body</p>")
    assert "Hello HTML" in text(run("html2pdf", [html])[0])
    pdfa = run("pdfa", [report])[0]
    assert pdfa.exists() and pages(pdfa) == 4


def test_rejects_wrong_files(tmp_path):
    with pytest.raises(ValueError, match="at least 2"):
        run("merge", [tmp_path / "one.pdf"])
    with pytest.raises(ValueError, match="can't use"):
        run("compress", [tmp_path / "photo.png"])
